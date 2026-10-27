"""Streamlit UI for Malaria SSR1.6 validation.
Run: streamlit run malaria_streamlit_app.py
Upload the trusted rule Python file in the UI; configure ADMIN_PASSWORD in Streamlit secrets.
"""

from __future__ import annotations

import importlib.util
import hashlib
import hmac
import tempfile
from io import BytesIO
from pathlib import Path

import pandas as pd
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

BASE_DIR = Path(__file__).resolve().parent
TEMPLATE_FILE = BASE_DIR / "Malaria_SSR1_6_Input_Template.xlsx"
LOCAL_PHOTO = BASE_DIR / "anopheles_mosquito.jpg"
PHOTO_URL = "https://commons.wikimedia.org/wiki/Special:FilePath/Anopheles_minimus.jpg?width=960"

st.set_page_config(page_title="Malaria SSR1.6 | Data Validation", page_icon="🦟", layout="wide")
st.markdown("""
<style>
.block-container {padding-top: 1.5rem;}
.hero-title {font-size:2rem; font-weight:750; color:#146451; margin-bottom:.3rem;}
.hero-sub {color:#586a68; font-size:1.0rem;}
</style>
""", unsafe_allow_html=True)


def load_uploaded_rules(rule_bytes: bytes):
    """Execute a trusted Python rules module from a temporary server-side file."""
    digest = hashlib.sha256(rule_bytes).hexdigest()[:16]
    with tempfile.TemporaryDirectory(prefix="malaria_rules_") as folder:
        rule_path = Path(folder) / "uploaded_malaria_rules.py"
        rule_path.write_bytes(rule_bytes)
        spec = importlib.util.spec_from_file_location(f"malaria_rules_{digest}", rule_path)
        if spec is None or spec.loader is None:
            raise RuntimeError("Unable to load uploaded rules file.")
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
    required = ("malaria_positive", "malaria_aggregate", "malaria_recheck")
    absent = [name for name in required if not callable(getattr(module, name, None))]
    if absent:
        raise ValueError("Rule script is missing function(s): " + ", ".join(absent))
    return module


def admin_authorized():
    """Require administrator password before accepting executable rule uploads."""
    password = st.secrets.get("ADMIN_PASSWORD", "")
    if not password:
        st.error("Admin access is not configured. Set ADMIN_PASSWORD in Streamlit Cloud → App settings → Secrets.")
        return False
    submitted = st.text_input("Administrator password", type="password", key="rule_admin_password")
    if not submitted:
        return False
    if not hmac.compare_digest(submitted, str(password)):
        st.error("Incorrect administrator password.")
        return False
    return True


def read_sheet(upload, sheet_name):
    # Keep data as supplied, without allowing inferred numeric conversion of IDs.
    # Native Excel dates remain as timestamps; the rules decide acceptable date formats.
    return pd.read_excel(upload, sheet_name=sheet_name, dtype=object, engine="openpyxl")


def select_sheet(upload, key):
    if upload is None:
        return None
    xl = pd.ExcelFile(upload, engine="openpyxl")
    return st.selectbox("Worksheet", xl.sheet_names, key=key)


def to_excel(positive=None, aggregate=None, recheck=None):
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        if positive is not None:
            positive.to_excel(writer, sheet_name="Positive_Checked", index=False)
        if aggregate is not None:
            aggregate.to_excel(writer, sheet_name="Aggregate_Checked", index=False)
        if recheck is not None:
            rows = recheck.get("custom_rows")
            if rows:
                wb = writer.book
                ws = wb.create_sheet(recheck.get("sheet_title", "REACH_RECHECK")[:31])
                for row in rows:
                    ws.append([None if pd.isna(v) and not isinstance(v, str) else v for v in row])
                for cell in ws[1]:
                    cell.fill = PatternFill("solid", fgColor="146451")
                    cell.font = Font(color="FFFFFF", bold=True)
                for cell in ws[2]:
                    cell.fill = PatternFill("solid", fgColor="DCECE8")
                    cell.font = Font(bold=True)
                    cell.alignment = Alignment(wrap_text=True, vertical="center")
                ws.freeze_panes = "G3"
            else:
                recheck["file"].to_excel(writer, sheet_name="REACH_RECHECK", index=False)
        for ws in writer.book.worksheets:
            if ws.title != "REACH_RECHECK":
                ws.freeze_panes = "A2"
                ws.auto_filter.ref = ws.dimensions
                for cell in ws[1]:
                    cell.fill = PatternFill("solid", fgColor="146451")
                    cell.font = Font(color="FFFFFF", bold=True)
                for column in ws.columns:
                    letter = column[0].column_letter
                    ws.column_dimensions[letter].width = min(35, max(13, len(str(column[0].value or "")) + 3))
    return output.getvalue()


def count_comments(df):
    if "COMMENT" not in df:
        return 0
    comments = df["COMMENT"].fillna("").astype(str).str.strip()
    return int(comments.ne("").sum())


heading, photo = st.columns([3.2, 1], vertical_alignment="center")
with heading:
    st.markdown('<div class="hero-title">Malaria SSR1.6 — Data Validation</div>', unsafe_allow_html=True)
    st.markdown('<div class="hero-sub">Upload your trusted Python rules and malaria Excel records, run checks, and download results.</div>', unsafe_allow_html=True)
with photo:
    # A local copy works offline; otherwise Streamlit loads the Wikimedia Commons photo.
    try:
        st.image(str(LOCAL_PHOTO) if LOCAL_PHOTO.exists() else PHOTO_URL,
                 caption="Anopheles minimus • Photo: CDC / James Gathany (public domain)",
                 use_container_width=True)
    except Exception:
        st.caption("Mosquito photo not available offline. Add anopheles_mosquito.jpg beside this app.")

st.divider()
if TEMPLATE_FILE.is_file():
    st.download_button("Download blank Excel input template", data=TEMPLATE_FILE.read_bytes(),
                       file_name=TEMPLATE_FILE.name,
                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

st.subheader("1. Upload malaria validation rules")
st.caption("Security: Python files run on the app server. Only upload rule scripts you trust. Administrator access is required.")
if not admin_authorized():
    st.stop()
rule_upload = st.file_uploader("Upload malaria rules (.py)", type=["py"], key="rule_file")
if rule_upload is not None:
    st.caption(f"Selected rules: {rule_upload.name}")

st.subheader("2. Upload Excel records")

mode = st.radio("Run checks", ["Positive + Aggregate", "Positive only", "Aggregate only"],
                horizontal=True, index=0)
need_pos = mode != "Aggregate only"
need_agg = mode != "Positive only"

left, right = st.columns(2)
with left:
    st.subheader("Positive cases")
    pos_upload = st.file_uploader("Upload positive records (.xlsx)", type=["xlsx"],
                                  key="positive", disabled=not need_pos)
    pos_sheet = select_sheet(pos_upload, "positive_sheet") if need_pos else None
with right:
    st.subheader("Aggregate / total records")
    agg_upload = st.file_uploader("Upload aggregate records (.xlsx)", type=["xlsx"],
                                  key="aggregate", disabled=not need_agg)
    agg_sheet = select_sheet(agg_upload, "aggregate_sheet") if need_agg else None

can_run = (rule_upload is not None and (not need_pos or pos_upload is not None) and (not need_agg or agg_upload is not None))
run = st.button("Run malaria validation", type="primary", disabled=not can_run,
                use_container_width=True)

if run:
    try:
        rules = load_uploaded_rules(rule_upload.getvalue())
        positive_source = read_sheet(pos_upload, pos_sheet) if need_pos else None
        aggregate_source = read_sheet(agg_upload, agg_sheet) if need_agg else None
        with st.spinner("Checking records against SSR1.6 rules..."):
            checked_pos = rules.malaria_positive(positive_source) if need_pos else None
            checked_agg = rules.malaria_aggregate(aggregate_source) if need_agg else None
            recheck = None
            if checked_pos is not None and checked_agg is not None:
                recheck = rules.malaria_recheck(
                    checked_pos, checked_agg,
                    original_input_columns=set(positive_source.columns),
                )
            result_file = to_excel(checked_pos, checked_agg, recheck)
        st.session_state["result"] = (checked_pos, checked_agg, recheck, result_file)
    except Exception as exc:
        st.session_state.pop("result", None)
    
        st.error(
            f"Validation failed: {type(exc).__name__}: {repr(exc)}"
        )
    
        with st.expander("Show detailed error"):
            st.exception(exc)

if "result" in st.session_state:
    checked_pos, checked_agg, recheck, result_file = st.session_state["result"]
    st.success("Validation finished. Review findings and download the checked workbook.")
    metrics = st.columns(4)
    metrics[0].metric("Positive records", len(checked_pos) if checked_pos is not None else "—")
    metrics[1].metric("Positive rows with comments", count_comments(checked_pos) if checked_pos is not None else "—")
    metrics[2].metric("Aggregate records", len(checked_agg) if checked_agg is not None else "—")
    metrics[3].metric("Aggregate rows with comments", count_comments(checked_agg) if checked_agg is not None else "—")

    st.download_button("Download all validation results (.xlsx)", data=result_file,
                       file_name="Malaria_SSR1_6_Validation_Results.xlsx",
                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                       type="primary")

    tabs = []
    if checked_pos is not None:
        tabs.append(("Positive results", checked_pos))
    if checked_agg is not None:
        tabs.append(("Aggregate results", checked_agg))
    if recheck is not None:
        tabs.append(("REACH_RECHECK preview", recheck.get("display", pd.DataFrame())))
    tab_widgets = st.tabs([x[0] for x in tabs])
    for tab, (_, df) in zip(tab_widgets, tabs):
        with tab:
            if "COMMENT" in df.columns:
                only_findings = st.checkbox("Show rows with comments only", key=f"filter_{tab_widgets.index(tab)}")
                if only_findings:
                    df = df[df["COMMENT"].fillna("").astype(str).str.strip().ne("")]
            st.dataframe(df, use_container_width=True, hide_index=True)
    if recheck is not None:
        st.info("REACH_RECHECK includes Excel formulas. Open the downloaded workbook in Excel or LibreOffice to recalculate the comparison fields.")
